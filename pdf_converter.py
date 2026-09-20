import io
import json
import logging
import os
import re
import uuid
from pathlib import Path
from typing import Tuple, Dict, Any

import pypdf
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import DATA_DIR, GEMINI_API_KEY
from google import genai
from google.genai import types

logger = logging.getLogger(__name__)

EXPORTS_DIR = DATA_DIR / "exports"
EXPORTS_DIR.mkdir(exist_ok=True)

def extract_text_from_pdf(doc_bytes: bytes) -> str:
    try:
        reader = pypdf.PdfReader(io.BytesIO(doc_bytes))
        pages = []
        for i, page in enumerate(reader.pages):
            txt = page.extract_text() or ""
            if txt.strip():
                pages.append(f"[Halaman {i+1}]:\n{txt.strip()}")
        return "\n\n".join(pages)
    except Exception as e:
        logger.error(f"Gagal ekstraksi teks PDF: {e}")
        return ""

def create_excel_from_structured_data(data: Dict[str, Any], output_path: str):
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Ringkasan Laporan"
    ws1.views.sheetView[0].showGridLines = True

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
    regular_font = Font(name="Calibri", size=11)
    bold_font = Font(name="Calibri", size=11, bold=True)
    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    ws1['A1'] = data.get('title', 'Laporan Hasil Konversi PDF')
    ws1['A1'].font = title_font

    ws1.cell(row=3, column=1, value="Parameter / Rincian").font = header_font
    ws1.cell(row=3, column=1).fill = header_fill
    ws1.cell(row=3, column=2, value="Nilai / Keterangan").font = header_font
    ws1.cell(row=3, column=2).fill = header_fill

    row = 4
    for item in data.get('summary', []):
        ws1.cell(row=row, column=1, value=str(item.get('label', ''))).font = bold_font
        c2 = ws1.cell(row=row, column=2, value=str(item.get('value', '')))
        c2.font = regular_font
        ws1.cell(row=row, column=1).border = thin_border
        ws1.cell(row=row, column=2).border = thin_border
        row += 1

    # Format kolom summary
    for col in ws1.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = max(max_len + 4, 16)

    # Tambahkan sheet tabel jika ada
    for table in data.get('tables', []):
        t_name = str(table.get('name', 'Data Detail'))[:30]
        ws = wb.create_sheet(title=t_name)
        ws.views.sheetView[0].showGridLines = True

        headers = table.get('headers', [])
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=str(h))
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        for r_idx, r_data in enumerate(table.get('rows', []), 2):
            for c_idx, val in enumerate(r_data, 1):
                c = ws.cell(row=r_idx, column=c_idx, value=str(val))
                c.font = regular_font
                c.border = thin_border

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(output_path)

async def convert_pdf_document_to_excel(doc_bytes: bytes, file_name: str, caption: str = "") -> Tuple[str, str]:
    client = genai.Client(api_key=GEMINI_API_KEY)
    extracted_text = extract_text_from_pdf(doc_bytes)

    prompt = (
        f"Anda adalah asisten ahli ekstraksi data dokumen dan pelaporan Excel.\n"
        f"Dokumen: {file_name}. Caption pengguna: '{caption}'\n"
        f"Tugas Anda:\n"
        f"1. Analisis seluruh teks dokumen ini secara mendalam.\n"
        f"2. Ekstrak data laporan menjadi struktur JSON yang valid dengan skema:\n"
        f"`json\n"
        f"{{\n"
        f'  "title": "Judul Laporan",\n'
        f'  "summary": [\n'
        f'    {{"label": "Nama Parameter", "value": "Nilai / Jumlah"}},\n'
        f'    ...\n'
        f'  ],\n'
        f'  "tables": [\n'
        f'    {{\n'
        f'      "name": "Nama Tabel",\n'
        f'      "headers": ["Kolom 1", "Kolom 2", ...],\n'
        f'      "rows": [["Nilai 1", "Nilai 2", ...], ...]\n'
        f'    }}\n'
        f'  ],\n'
        f'  "executive_summary": "Rangkuman singkat laporan untuk pesan chat WhatsApp/Telegram"\n'
        f"}}\n"
        f"`\n"
        f"PENTING: Keluarkan HANYA blok JSON di dalam `json ... ` agar bisa diparsing otomatis.\n\n"
        f"TEKS DOKUMEN:\n{extracted_text[:18000]}"
    )

    models_to_try = [
        "gemini-3.1-flash-lite",
        "gemini-flash-lite-latest",
        "gemini-3.5-flash-lite",
        "gemini-flash-latest"
    ]

    response_text = ""
    for m in models_to_try:
        try:
            resp = client.models.generate_content(
                model=m,
                contents=prompt,
                config=types.GenerateContentConfig(temperature=0.2)
            )
            response_text = resp.text or ""
            if response_text:
                break
        except Exception as e:
            logger.warning(f"Model {m} gagal saat ekstraksi PDF ke Excel: {e}")
            continue

    # Parsing JSON dari respon Gemini
    json_match = re.search(r'`json\s*([\s\S]*?)\s*`', response_text)
    if json_match:
        raw_json = json_match.group(1)
    else:
        raw_json = response_text.strip()

    try:
        parsed_data = json.loads(raw_json)
    except Exception:
        # Fallback jika parsing JSON gagal
        parsed_data = {
            "title": f"Laporan {Path(file_name).stem}",
            "summary": [{"label": "Status", "value": "Teks berhasil diekstrak dari PDF"}],
            "tables": [{
                "name": "Isi Dokumen",
                "headers": ["Halaman / Cuplikan"],
                "rows": [[line] for line in extracted_text.splitlines() if line.strip()][:50]
            }],
            "executive_summary": response_text[:500] if response_text else "Laporan berhasil diekstrak ke Excel."
        }

    clean_name = re.sub(r'[^a-zA-Z0-9_-]', '_', Path(file_name).stem)
    out_file = EXPORTS_DIR / f"Laporan_{clean_name}_{uuid.uuid4().hex[:6]}.xlsx"
    create_excel_from_structured_data(parsed_data, str(out_file))

    summary_text = parsed_data.get('executive_summary', '')
    if not summary_text or len(summary_text) < 10:
        summary_text = (
            f"📄 *LAPORAN TELAH DIUBAH KE EXCEL*\n\n"
            f"• **Judul Dokumen:** {parsed_data.get('title', file_name)}\n"
            f"• **Status:** File spreadsheet Excel (.xlsx) berhasil dibuat khusus dari PDF ini tanpa mengubah database harian Anda."
        )

    return summary_text, str(out_file)
