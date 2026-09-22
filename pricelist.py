import csv
import re
import difflib
import io
import logging
from pathlib import Path
from typing import Dict, Any, List, Optional, Tuple

logger = logging.getLogger(__name__)

PRICELIST_PATH = Path("data/pricelist.csv")
PRICELIST_FIELDNAMES = [
    "Model", "Description", "Harga_MD", "Harga_IPP", "Harga_SDP",
    "Harga_ADP", "Harga_DPP", "Harga_MDP", "Harga_MSRP", "Harga_Non_DPP",
    "Category", "Sumber", "Warranty", "Brand"
]

def infer_brand(model: str) -> str:
    """Menebak merek produk berdasarkan kode atau awalan model."""
    m = model.strip().upper()
    if m.startswith("DH-") or m.startswith("DHI-") or m.startswith("DH") or m.startswith("DHI"):
        return "Dahua"
    if m.startswith("DS-") or m.startswith("IDS-") or m.startswith("HC-") or m.startswith("HIK"):
        return "Hikvision"
    if m.startswith("RG-") or m.startswith("REYEE") or m.startswith("ES-") or m.startswith("NBS-") or m.startswith("RAP"):
        return "Ruijie"
    if m.startswith("THC-") or m.startswith("HL-") or m.startswith("HILOOK"):
        return "HiLook"
    if m.startswith("HV-") or m.startswith("HIVIEW"):
        return "Hiview"
    if m.startswith("IPC-") or m.startswith("IMOU"):
        return "Imou"
    return "Umum"

def normalize_code(code: str) -> str:
    """Bersihkan kode: lowercase, hapus spasi, tanda hubung, dan tanda baca."""
    if not code:
        return ""
    cleaned = re.sub(r'[^a-zA-Z0-9]', '', code).lower()
    return cleaned

def load_pricelist() -> List[Dict[str, Any]]:
    """Muat seluruh data pricelist dari file CSV."""
    if not PRICELIST_PATH.exists():
        return []
    
    products = []
    with open(PRICELIST_PATH, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for row in reader:
            products.append(row)
    return products

def format_rupiah_num(val: Any) -> str:
    """Format angka integer/float ke format Rupiah standar: Rp 1.250.000."""
    try:
        clean_val = re.sub(r'[^0-9]', '', str(val))
        if not clean_val:
            return "Rp 0"
        num = int(clean_val)
        return f"Rp {num:,}".replace(",", ".")
    except Exception:
        return f"Rp {val}"

def search_pricelist(query: str, price_tier: str = "", brand_filter: str = "") -> Dict[str, Any]:
    """
    Cari produk di pricelist berdasarkan kode, model, atau nama.
    Menerapkan Aturan 1, 2, 3, 4, 5.
    """
    products = load_pricelist()
    if not products:
        return {
            "status": "empty",
            "message": "Data pricelist belum tersedia di database."
        }
    
    q_raw = query.strip()
    
    # Deteksi merek dalam query jika ada (Aturan 3)
    detected_brand = brand_filter.strip().lower() if brand_filter else ""
    brands = ["hikvision", "dahua", "ruijie", "hilook", "hiview", "imou", "reyee"]
    if not detected_brand:
        for b in brands:
            pattern = rf'\b{b}\b'
            if re.search(pattern, q_raw, flags=re.IGNORECASE):
                detected_brand = "ruijie" if b == "reyee" else b
                q_raw = re.sub(pattern, '', q_raw, flags=re.IGNORECASE).strip()
                break
                
    if detected_brand:
        products = [p for p in products if p.get("Brand", "").lower() == detected_brand]
        if not products:
            return {
                "status": "not_found",
                "query": query,
                "suggestions": []
            }

    q_norm = normalize_code(q_raw)

    # Deteksi pencarian kategori umum (misal: "home router", "router wifi", "home router wifi")
    if q_norm in [
        "homerouter", "homerouterwifi", "routerwifi", "wifirouter",
        "routerhome", "reyerouter", "ruijierouter", "routerruijie",
        "homerouterreyee", "daftarmodelrouter", "homerouterruijie", "router"
    ]:
        router_matches = [p for p in products if p.get("Category") == "Home Router Wi-Fi" or "Home Router" in p.get("Category", "")]
        if router_matches:
            return {
                "status": "ambiguous",
                "matches": router_matches,
                "price_tier": price_tier
            }

    # 1. Cari exact match (persis sama setelah dinormalisasi)
    exact_matches = [p for p in products if q_norm == normalize_code(p.get("Model", ""))]
    if len(exact_matches) == 1:
        return {
            "status": "exact",
            "product": exact_matches[0],
            "price_tier": price_tier
        }
    elif len(exact_matches) > 1:
        return {
            "status": "ambiguous",
            "matches": exact_matches,
            "price_tier": price_tier
        }
    
    # 2. Coba tanpa awalan brand jika ada ('dhi', 'dh', 'ids', 'ds', 'rg', 'reyee', 'hilook', 'hiview', 'imou', 'thc', 'ipc', 'th', 'hv', 'hik')
    def strip_brand_pfx(s: str) -> str:
        for pfx in ['reyee', 'dhi', 'dh', 'ids', 'ds', 'rg', 'hilook', 'hiview', 'imou', 'thc', 'ipc', 'th', 'hv', 'hik']:
            if s.startswith(pfx):
                return s[len(pfx):]
        return s

    q_stripped = strip_brand_pfx(q_norm)
    pfx_matches = []
    for p in products:
        m_norm = normalize_code(p.get("Model", ""))
        m_stripped = strip_brand_pfx(m_norm)
        if q_stripped == m_stripped:
            pfx_matches.append(p)
            
    if len(pfx_matches) == 1:
        return {
            "status": "exact",
            "product": pfx_matches[0],
            "price_tier": price_tier
        }
    elif len(pfx_matches) > 1:
        return {
            "status": "ambiguous",
            "matches": pfx_matches,
            "price_tier": price_tier
        }

    # 3. Cari matches awalan / substring
    substring_matches = []
    for p in products:
        m_norm = normalize_code(p.get("Model", ""))
        m_stripped = strip_brand_pfx(m_norm)
        if q_norm in m_norm or (len(q_stripped) >= 3 and q_stripped in m_stripped):
            substring_matches.append(p)
            
    if len(substring_matches) == 1:
        return {
            "status": "exact",
            "product": substring_matches[0],
            "price_tier": price_tier
        }
    elif len(substring_matches) > 1 and len(substring_matches) <= 6:
        return {
            "status": "ambiguous",
            "matches": substring_matches,
            "price_tier": price_tier
        }
    
    # 4. Jika tidak ditemukan, gunakan similarity (difflib) untuk 3 kandidat paling mirip
    all_models = [p.get("Model", "") for p in products]
    all_norms = {normalize_code(m): m for m in all_models}
    close_norms = difflib.get_close_matches(q_norm, list(all_norms.keys()), n=3, cutoff=0.4)
    if not close_norms:
        # Coba juga pencarian kata kunci pada Description
        desc_matches = []
        words = [w for w in q_raw.lower().split() if len(w) > 2]
        for p in products:
            desc_lower = p.get("Description", "").lower()
            if any(w in desc_lower for w in words):
                desc_matches.append(p)
                if len(desc_matches) >= 3:
                    break
        if desc_matches:
            return {
                "status": "not_found_with_suggestions",
                "query": q_raw,
                "suggestions": [p.get("Model", "") for p in desc_matches]
            }

        return {
            "status": "not_found",
            "query": q_raw,
            "suggestions": []
        }
    
    similar_models = [all_norms[k] for k in close_norms]
    return {
        "status": "not_found_with_suggestions",
        "query": q_raw,
        "suggestions": similar_models[:3]
    }

def _get_price(product: Dict[str, Any], *keys: str) -> Optional[str]:
    """Mengambil harga pertama yang tersedia dan valid (bukan 0 atau kosong)."""
    for k in keys:
        v = product.get(k)
        if v is not None and str(v).strip() not in ("", "0", "/"):
            return str(v).strip()
    return None

def format_single_product_answer(product: Dict[str, Any], tier: str = "") -> str:
    """Format: Model dan Harga saja untuk semua tipe dan merek."""
    model = product.get("Model", "N/A")
    brand = product.get("Brand", "General")
    tier_upper = tier.upper().strip() if tier else ""
    
    brand_lower = brand.lower()
    if brand_lower == "hikvision":
        price_val = _get_price(product, "Harga_DPP", "Harga_ADP", "Harga_MD", "Harga_MSRP")
    elif brand_lower == "dahua":
        # Wajib utamakan harga MD/MDP, tapi kalau tidak ada harga MD dan cuma ada MSRP, berikan MSRP
        price_val = _get_price(product, "Harga_MDP", "Harga_MD", "Harga_ADP", "Harga_DPP", "Harga_MSRP")
    elif brand_lower == "ruijie":
        price_val = _get_price(product, "Harga_ADP", "Harga_MD", "Harga_MSRP")
    elif brand_lower == "hilook":
        price_val = _get_price(product, "Harga_MD", "Harga_ADP", "Harga_MSRP")
    elif brand_lower == "hiview":
        price_val = _get_price(product, "Harga_MD", "Harga_ADP", "Harga_MSRP")
    elif brand_lower == "imou":
        if tier_upper == "SDP":
            price_val = _get_price(product, "Harga_SDP", "Harga_ADP", "Harga_MSRP")
        elif tier_upper in ("SRP", "MSRP"):
            price_val = _get_price(product, "Harga_MSRP", "Harga_SRP")
        elif tier_upper == "ONLINE":
            price_val = _get_price(product, "Harga_Non_DPP", "Harga_Online", "Harga_MSRP")
        else:
            # Default IPP
            price_val = _get_price(product, "Harga_IPP", "Harga_MD", "Harga_ADP", "Harga_MSRP")
    else:
        price_val = _get_price(product, "Harga_ADP", "Harga_MD", "Harga_MSRP")

    if tier_upper == "MSRP":
        msrp_candidate = _get_price(product, "Harga_MSRP")
        if msrp_candidate:
            price_val = msrp_candidate
    elif tier_upper in ("NON-DPP", "NONDPP", "NON-MD", "NONMD"):
        nondpp_candidate = _get_price(product, "Harga_Non_DPP")
        if nondpp_candidate:
            price_val = nondpp_candidate

    if not price_val:
        formatted_price = "harga tidak tersedia"
    else:
        formatted_price = format_rupiah_num(price_val)

    return f"{model} : {formatted_price}"

def query_pricelist_tool(query: str, tier: str = "") -> str:
    """Fungsi pembantu yang dipanggil oleh Gemini Agent atau command bot.
    Mendukung pencarian 1 tipe maupun sekaligus banyak tipe.
    Format ringkas: Model dan Harga saja.
    """
    cleaned = query.replace(";", "\n").replace(",", "\n")
    cleaned = re.sub(r'\s+(?:dan|&)\s+', '\n', cleaned, flags=re.IGNORECASE)
    parts = [p.strip() for p in cleaned.split("\n") if p.strip()]
    
    if len(parts) > 1:
        results = []
        for part in parts:
            p_clean = re.sub(r'^(?:tolong\s+)?(?:carikan\s+)?(?:harga\s+)?(?:adp\s+)?(?:md\s+)?(?:dpp\s+)?(?:dealer\s+)?(?:untuk\s+)?', '', part, flags=re.IGNORECASE).strip()
            p_clean = re.sub(r'\s+(?:berapa|dong|ya|unit|pcs)$', '', p_clean, flags=re.IGNORECASE).strip()
            if not p_clean:
                continue
            res = search_pricelist(p_clean, tier)
            st = res.get("status")
            if st == "exact":
                results.append(format_single_product_answer(res["product"], tier))
            elif st == "ambiguous":
                m_list = [format_single_product_answer(m, tier) for m in res['matches'][:5]]
                results.append("\n".join(m_list))
            elif st == "not_found_with_suggestions" and res.get("suggestions"):
                sugs = [f"• {s}" for s in res["suggestions"]]
                results.append(f"Model {p_clean} tidak ada di pricelist. Mungkin yang mirip:\n" + "\n".join(sugs))
            else:
                results.append(f"Model {p_clean} tidak ada di pricelist.")
        return "\n".join(results)
    
    # 1 tipe saja
    result = search_pricelist(query, tier)
    st = result.get("status")
    
    if st == "exact":
        return format_single_product_answer(result["product"], tier)
        
    elif st == "ambiguous":
        items = result["matches"]
        lines = [format_single_product_answer(p, tier) for p in items[:15]]
        return "\n".join(lines)
        
    elif st == "not_found_with_suggestions":
        sug = result.get("suggestions", [])
        if sug:
            sugs = [f"• {s}" for s in sug]
            return f"Model {query} tidak ada di pricelist. Mungkin yang mirip:\n" + "\n".join(sugs)
        return f"Model {query} tidak ada di pricelist."
        
    else:
        return f"Model {query} tidak ada di pricelist."

def try_direct_pricelist_query(text: str) -> Optional[str]:
    """
    Mendeteksi apakah teks pengguna adalah murni kueri model/tipe produk.
    Jika ya, langsung kembalikan format jawaban harga tanpa melalui LLM.
    """
    t = text.strip()
    if not t or t.startswith('/'):
        return None

    lower = t.lower()
    conv_words = [
        'catat', 'beli', 'makan', 'pengeluaran', 'pemasukan', 'saldo', 'reset',
        'laporan', 'excel', 'jadwal', 'ingatkan', 'alarm', 'tugas', 'todo',
        'halo', 'hai', 'pagi', 'siang', 'malam', 'siapa', 'kamu', 'bisa apa',
        'kenapa', 'gimana', 'bagaimana', 'diskon', 'total', 'kalau', 'unit',
        'pcs', 'ubah', 'ganti', 'set', 'update', 'perbaiki', 'benerin', 'tolong'
    ]
    words_in_text = lower.split()
    if any(w in words_in_text for w in conv_words):
        return None

    # Hapus awalan cek harga jika ada
    candidate = re.sub(r'^(?:cek\s+harga|berapa\s+harga|harga|pricelist)\s+', '', t, flags=re.IGNORECASE).strip()
    if not candidate:
        return None

    # Pisahkan bagian dengan koma, titik koma, atau newline
    raw_parts = [p.strip() for p in re.split(r'[,;\n]+', candidate) if p.strip()]
    if not raw_parts or len(raw_parts) > 30:
        return None

    # Setiap bagian harus mirip kode model produk
    for p in raw_parts:
        if len(p) < 2 or len(p) > 40:
            return None
        if not re.match(r'^[a-zA-Z0-9\-\(\)\/\.\_\+\s]+$', p):
            return None
        # Harus mengandung angka atau awalan merek resmi
        p_low = p.lower()
        has_digit = bool(re.search(r'[0-9]', p))
        has_brand_pfx = bool(re.search(r'^(?:rg|reyee|dh|dhi|ds|ids|hc|thc|hilook|imou|hiview|hv|es|nbs)', p_low))
        if not (has_digit or has_brand_pfx):
            return None

    res = query_pricelist_tool(candidate)
    # Jika semua bagian dinyatakan tidak ada di pricelist, serahkan ke LLM/Gemini
    if 'tidak ada di pricelist' in res:
        all_not_found = all(f'Model {p} tidak ada di pricelist' in res for p in raw_parts)
        if all_not_found:
            return None

    return res

def save_new_pricelist_csv(content: str) -> Tuple[bool, str, int]:
    """Menyimpan file CSV pricelist baru yang diupload user."""
    try:
        lines = [l for l in content.strip().splitlines() if l.strip()]
        if len(lines) < 2:
            return False, "File CSV kosong atau tidak memiliki baris data.", 0
        
        header = lines[0].lower()
        if "model" not in header and "tipe" not in header:
            return False, "Format CSV harus memiliki kolom 'Model' atau 'Tipe'.", 0
            
        with open(PRICELIST_PATH, "w", encoding="utf-8") as f:
            f.write(content.strip())
            
        products = load_pricelist()
        return True, f"Pricelist berhasil diperbarui dengan {len(products)} produk.", len(products)
    except Exception as e:
        return False, f"Gagal menyimpan pricelist CSV: {e}", 0

def extract_text_from_excel(doc_bytes: bytes, max_rows: int = 150) -> str:
    """Ekstrak isi teks/tabel dari file Excel (.xlsx) untuk dianalisis oleh AI."""
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(doc_bytes), data_only=True)
        sheets_data = []
        for name in wb.sheetnames[:5]:
            ws = wb[name]
            lines = [f"=== Sheet: {name} ==="]
            rows = list(ws.iter_rows(values_only=True))
            for r in rows[:max_rows]:
                if not any(r):
                    continue
                r_str = [str(c) if c is not None else "" for c in r]
                lines.append(" | ".join(r_str))
            if len(rows) > max_rows:
                lines.append(f"... (dan {len(rows) - max_rows} baris lainnya)")
            sheets_data.append("\n".join(lines))
        return "\n\n".join(sheets_data)
    except Exception as e:
        logger.error(f"Gagal ekstrak excel: {e}")
        return ""

def import_pricelist_from_excel(doc_bytes: bytes, file_name: str = "") -> Tuple[bool, str, int]:
    """Membaca file Excel (.xlsx) semua sheet dan menggabungkan (merge) ke data/pricelist.csv tanpa menghapus data produk yang sudah ada."""
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(doc_bytes), data_only=True)
        new_extracted = {}

        def clean_num(val):
            if not val:
                return ""
            try:
                if isinstance(val, (int, float)):
                    v = int(round(val))
                    return str(v) if v > 0 else ""
            except Exception:
                pass
            s = str(val).strip()
            if s.lower() in ("none", "n/a", "-", "null", "tbd", "na", "--"):
                return ""
            c = re.sub(r'[^0-9]', '', s)
            return c if c and int(c) > 0 else ""

        for sname in wb.sheetnames:
            sname_clean = sname.strip()
            if sname_clean.lower() in ["update", "history", "changelog", "hot model list for retail", "hot model list for smb project"]:
                continue
            ws = wb[sname]
            rows = list(ws.iter_rows(values_only=True))
            if not rows or len(rows) < 2:
                continue

            # Kasus khusus: Hik-Connect Team
            if "hik-connect team" in sname_clean.lower():
                for r in rows[2:]:
                    if len(r) > 4 and r[4]:
                        name = str(r[4]).strip()
                        if name.startswith("HC-"):
                            dpp = clean_num(r[8]) if len(r) > 8 else ""
                            msrp = clean_num(r[7]) if len(r) > 7 else ""
                            nondpp = clean_num(r[9]) if len(r) > 9 else ""
                            desc = str(r[5]).strip() if len(r) > 5 and r[5] else ""
                            final_dpp = dpp or nondpp or msrp
                            if final_dpp:
                                new_extracted[name] = {
                                    "Model": name,
                                    "Description": re.sub(r'\s+', ' ', desc)[:200],
                                    "Harga_MD": final_dpp,
                                    "Harga_IPP": final_dpp,
                                    "Harga_SDP": final_dpp,
                                    "Harga_ADP": final_dpp,
                                    "Harga_DPP": final_dpp,
                                    "Harga_MDP": final_dpp,
                                    "Harga_MSRP": msrp or final_dpp,
                                    "Harga_Non_DPP": nondpp,
                                    "Category": "Hik-Connect Team",
                                    "Sumber": "Hikvision Hik-Connect Team",
                                    "Warranty": "2 Years Warranty",
                                    "Brand": "Hikvision"
                                }
                continue

            # Kasus khusus: Hik-Partner Pro
            if "hik-partner pro" in sname_clean.lower():
                for r in rows[3:]:
                    if len(r) > 1 and r[1]:
                        name = str(r[1]).strip()
                        if name.startswith("HPP-") or "Co-Branding" in name:
                            msrp = clean_num(r[5]) if len(r) > 5 else ""
                            dpp = clean_num(r[6]) if len(r) > 6 else ""
                            nondpp = clean_num(r[7]) if len(r) > 7 else ""
                            desc = str(r[3]).strip() if len(r) > 3 and r[3] else ""
                            final_dpp = dpp or nondpp or msrp
                            if final_dpp:
                                new_extracted[name] = {
                                    "Model": name,
                                    "Description": re.sub(r'\s+', ' ', desc)[:200],
                                    "Harga_MD": final_dpp,
                                    "Harga_IPP": final_dpp,
                                    "Harga_SDP": final_dpp,
                                    "Harga_ADP": final_dpp,
                                    "Harga_DPP": final_dpp,
                                    "Harga_MDP": final_dpp,
                                    "Harga_MSRP": msrp or final_dpp,
                                    "Harga_Non_DPP": nondpp,
                                    "Category": "Hik-Partner Pro",
                                    "Sumber": "Hikvision Hik-Partner Pro",
                                    "Warranty": "2 Years Warranty",
                                    "Brand": "Hikvision"
                                }
                continue

            header_idx = -1
            for idx, r in enumerate(rows[:15]):
                r_str = [str(cell or "").lower().strip() for cell in r]
                if any("model" in c or "tipe" in c or "type" in c or "dpp" in c or (c == "name" and any("price" in x or "dpp" in x for x in r_str)) for c in r_str):
                    header_idx = idx
                    break

            if header_idx == -1:
                continue

            header_row = [str(cell or "").lower().strip() for cell in rows[header_idx]]

            col_order = -1
            col_basic = -1
            col_model = -1
            col_desc = -1
            col_dpp = -1
            col_nondpp = -1
            col_adp = -1
            col_md = -1
            col_msrp = -1
            col_warranty = -1

            for i, col_name in enumerate(header_row):
                if "order" in col_name and ("model" in col_name or "name" in col_name):
                    col_order = i
                elif "basic" in col_name and ("model" in col_name or "name" in col_name):
                    col_basic = i
                elif ("model" in col_name or "tipe" in col_name or "type" in col_name or col_name == "name") and col_model == -1:
                    col_model = i
                elif any(k in col_name for k in ["desc", "keterangan", "deskripsi", "spesifikasi", "feature"]):
                    if col_desc == -1:
                        col_desc = i
                elif "non-dpp" in col_name or "non dpp" in col_name or "nondpp" in col_name:
                    col_nondpp = i
                elif "dpp" in col_name and "non" not in col_name:
                    col_dpp = i
                elif "adp" in col_name:
                    col_adp = i
                elif "mdp" in col_name:
                    col_md = i
                elif "bottom" in col_name or "md" in col_name or "dealer" in col_name:
                    col_md = i
                elif any(k in col_name for k in ["msrp", "srp", "retail", "price list"]):
                    if col_msrp == -1 and "dpp" not in col_name:
                        col_msrp = i
                elif "warranty" in col_name or "garansi" in col_name:
                    col_warranty = i

            if col_dpp == -1 and col_adp == -1 and col_md == -1:
                for i, col_name in enumerate(header_row):
                    if "price" in col_name or "harga" in col_name:
                        col_dpp = i
                        break

            last_basic = ""
            last_desc = ""

            for r in rows[header_idx + 1:]:
                if not r or not any(r):
                    continue

                curr_basic = str(r[col_basic]).strip() if col_basic != -1 and len(r) > col_basic and r[col_basic] is not None else ""
                if curr_basic and curr_basic.lower() not in ("none", "null", "-", "standard"):
                    last_basic = curr_basic

                curr_desc = str(r[col_desc]).strip() if col_desc != -1 and len(r) > col_desc and r[col_desc] is not None else ""
                if curr_desc and curr_desc.lower() not in ("none", "null", "-"):
                    last_desc = curr_desc

                raw_order = str(r[col_order]).strip() if col_order != -1 and len(r) > col_order and r[col_order] else ""
                raw_model = str(r[col_model]).strip() if col_model != -1 and len(r) > col_model and r[col_model] else ""

                candidates = []
                for rm in [raw_order, raw_model]:
                    if not rm or rm.lower() in ("none", "model", "tipe", "type", "order model", "basic model", "name", "null", "-"):
                        continue
                    cm = re.sub(r'[\（\(](?:hot\s*sku|project|new|standard)[^\）\)]*[\）\)]', '', rm, flags=re.IGNORECASE).strip()
                    cm = cm.split('\n')[0].strip()
                    if len(cm) >= 3 and cm not in candidates:
                        candidates.append(cm)

                if not candidates and last_basic and len(last_basic) >= 3:
                    cm = re.sub(r'[\（\(](?:hot\s*sku|project|new|standard)[^\）\)]*[\）\)]', '', last_basic, flags=re.IGNORECASE).strip()
                    cm = cm.split('\n')[0].strip()
                    if cm and cm not in candidates:
                        candidates.append(cm)

                if not candidates:
                    continue

                # Ambil harga
                p_dpp = clean_num(r[col_dpp]) if col_dpp != -1 and len(r) > col_dpp else ""
                p_nondpp = clean_num(r[col_nondpp]) if col_nondpp != -1 and len(r) > col_nondpp else ""
                p_adp = clean_num(r[col_adp]) if col_adp != -1 and len(r) > col_adp else ""
                p_md = clean_num(r[col_md]) if col_md != -1 and len(r) > col_md else ""
                
                p_final = p_dpp or p_adp or p_md or p_nondpp
                if not p_final or int(p_final) == 0:
                    continue

                desc = curr_desc or last_desc
                desc = re.sub(r'\s+', ' ', desc).strip()[:200]
                msrp = clean_num(r[col_msrp]) if col_msrp != -1 and len(r) > col_msrp else p_final
                warranty = str(r[col_warranty]).strip() if col_warranty != -1 and len(r) > col_warranty and r[col_warranty] else "2 Years Warranty"

                # Deteksi Brand otomatis
                brand = "General"
                first_cand = candidates[0].upper()
                file_str = (file_name or "").lower()
                if "hilook" in file_str or first_cand.startswith("THC-") or first_cand.startswith("IPC-B12") or first_cand.startswith("IPC-D12") or first_cand.startswith("IPC-B14") or first_cand.startswith("IPC-D14") or first_cand.startswith("DVR-2") or first_cand.startswith("NVR-1"):
                    brand = "HiLook"
                elif "hiview" in file_str or first_cand.startswith("HV-") or first_cand.startswith("TH-") or first_cand.startswith("T1A20") or first_cand.startswith("B1A20") or first_cand.startswith("T1290") or first_cand.startswith("B1290"):
                    brand = "Hiview"
                elif "hik" in file_str or first_cand.startswith("DS-") or first_cand.startswith("IDS-") or first_cand.startswith("HC-") or first_cand.startswith("HPP-"):
                    brand = "Hikvision"
                elif "dahua" in file_str or first_cand.startswith("DH-") or first_cand.startswith("DHI-") or first_cand.startswith("XVR") or first_cand.startswith("NVR") or first_cand.startswith("HAC"):
                    brand = "Dahua"
                elif "ruijie" in file_str or "reyee" in file_str or first_cand.startswith("RG-") or first_cand.startswith("REYEE") or "RUIJIE" in first_cand:
                    brand = "Ruijie"

                for cand in candidates:
                    if cand.lower() in ("outdoor", "indoor", "waterproof", "analog", "bullet", "dome", "turret", "ptz", "accessories", "switch"):
                        continue
                    if brand == "Dahua":
                        dahua_md = p_md or p_dpp
                        new_extracted[cand] = {
                            "Model": cand,
                            "Description": desc,
                            "Harga_MD": dahua_md,
                            "Harga_IPP": dahua_md,
                            "Harga_SDP": p_dpp or dahua_md,
                            "Harga_ADP": dahua_md,
                            "Harga_DPP": p_dpp,
                            "Harga_MDP": dahua_md,
                            "Harga_MSRP": msrp,
                            "Harga_Non_DPP": p_nondpp or "",
                            "Category": sname_clean,
                            "Sumber": f"{brand} {sname_clean}",
                            "Warranty": warranty,
                            "Brand": brand
                        }
                    else:
                        new_extracted[cand] = {
                            "Model": cand,
                            "Description": desc,
                            "Harga_MD": p_md or p_final,
                            "Harga_IPP": p_final,
                            "Harga_SDP": p_final,
                            "Harga_ADP": p_adp or p_final,
                            "Harga_DPP": p_dpp or p_final,
                            "Harga_MDP": p_final,
                            "Harga_MSRP": msrp or p_final,
                            "Harga_Non_DPP": p_nondpp or "",
                            "Category": sname_clean,
                            "Sumber": f"{brand} {sname_clean}",
                            "Warranty": warranty,
                            "Brand": brand
                        }

        if not new_extracted:
            return False, "Tidak ada data produk yang berhasil diekstrak dari seluruh sheet.", 0

        # Muat produk yang sudah ada agar tidak terhapus (Merge)
        existing = {}
        if PRICELIST_PATH.exists():
            with open(PRICELIST_PATH, "r", encoding="utf-8") as f:
                reader = csv.DictReader(f)
                for r in reader:
                    m = r.get("Model", "").strip()
                    if m:
                        existing[m] = r

        # Gabungkan
        for m, item in new_extracted.items():
            existing[m] = item

        fieldnames = [
            "Model", "Description", "Harga_MD", "Harga_IPP", "Harga_SDP",
            "Harga_ADP", "Harga_DPP", "Harga_MDP", "Harga_MSRP", "Harga_Non_DPP",
            "Category", "Sumber", "Warranty", "Brand"
        ]
        with open(PRICELIST_PATH, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            for r in existing.values():
                writer.writerow({k: r.get(k, "") for k in fieldnames})

        return True, f"Berhasil mengimpor {len(new_extracted)} produk ke database (Total sekarang: {len(existing)} produk).", len(new_extracted)
    except Exception as e:
        logger.error(f"Gagal membaca file Excel: {e}")
        return False, f"Gagal membaca file Excel: {e}", 0

def update_product_price(model_query: str, new_price: Any, create_if_missing: bool = True) -> Tuple[bool, str]:
    """Mengubah atau memperbarui harga suatu produk di data/pricelist.csv."""
    clean_price = re.sub(r'[^0-9]', '', str(new_price))
    if not clean_price:
        return False, "Nominal harga tidak valid. Masukkan angka nominal yang benar, contoh: 450000 atau Rp 450.000."

    products = load_pricelist()
    q_norm = normalize_code(model_query)
    if not q_norm:
        return False, "Model produk tidak boleh kosong."

    found_idx = -1
    # 1. Exact match
    for idx, p in enumerate(products):
        m_norm = normalize_code(p.get("Model", ""))
        if q_norm == m_norm:
            found_idx = idx
            break

    # 2. Match without brand prefix
    if found_idx == -1:
        for idx, p in enumerate(products):
            m_norm = normalize_code(p.get("Model", ""))
            if (q_norm.startswith("rg") and q_norm[2:] == m_norm[2:]) or (q_norm.startswith("dh") and q_norm[2:] == m_norm[2:]):
                found_idx = idx
                break

    # 3. Substring match
    if found_idx == -1:
        for idx, p in enumerate(products):
            m_norm = normalize_code(p.get("Model", ""))
            if q_norm in m_norm or (len(q_norm) >= 5 and m_norm in q_norm):
                found_idx = idx
                break

    if found_idx != -1:
        p = products[found_idx]
        p["Harga_MD"] = clean_price
        p["Harga_ADP"] = clean_price
        p["Harga_MDP"] = clean_price
        p["Harga_DPP"] = clean_price
        p["Harga_IPP"] = clean_price
        p["Harga_SDP"] = clean_price
        updated_model = p.get("Model", model_query)
        action_desc = f"✅ Harga {updated_model} berhasil diupdate menjadi {format_rupiah_num(clean_price)}."
    elif create_if_missing:
        brand = infer_brand(model_query)
        new_prod = {k: "" for k in PRICELIST_FIELDNAMES}
        new_prod["Model"] = model_query.strip()
        new_prod["Brand"] = brand
        new_prod["Description"] = f"{brand} {model_query.strip()}"
        new_prod["Harga_MD"] = clean_price
        new_prod["Harga_ADP"] = clean_price
        new_prod["Harga_MDP"] = clean_price
        new_prod["Harga_DPP"] = clean_price
        new_prod["Harga_IPP"] = clean_price
        new_prod["Harga_SDP"] = clean_price
        new_prod["Category"] = "Manual Update"
        new_prod["Sumber"] = "User Direct Update"
        new_prod["Warranty"] = "Official Warranty"
        products.append(new_prod)
        updated_model = model_query.strip()
        action_desc = f"✅ Produk baru {updated_model} ({brand}) berhasil ditambahkan ke database dengan harga {format_rupiah_num(clean_price)}."
    else:
        return False, f"Produk '{model_query}' tidak ditemukan di database pricelist."

    with open(PRICELIST_PATH, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=PRICELIST_FIELDNAMES)
        writer.writeheader()
        for r in products:
            writer.writerow({k: r.get(k, "") for k in PRICELIST_FIELDNAMES})

    return True, action_desc

def batch_update_product_prices(input_text: str) -> Tuple[bool, str, int]:
    """
    Memperbarui banyak harga produk sekaligus dari teks multiline.
    Mendukung format:
    Model Rp100.000
    Model 100000
    Model : 100.000
    Model = 100.000
    """
    lines = [l.strip() for l in input_text.strip().splitlines() if l.strip()]
    if not lines:
        return False, "Teks pembaruan harga kosong.", 0

    products = load_pricelist()
    prod_map = {normalize_code(p.get("Model", "")): p for p in products}

    updated_count = 0
    added_count = 0
    results_detail = []

    for line in lines:
        line_clean = re.sub(r'^(?:/setharga|/updateharga|/ubahharga|/gantiharga|update harga|ubah harga|ganti harga|set harga|perbaiki harga)[:\s]*', '', line, flags=re.IGNORECASE).strip()
        if not line_clean:
            continue

        m = re.search(r'^(.*?)(?::\s*|=\s*|\s+(?:jadi|menjadi)\s*|\s+)(?:Rp\.?\s*)?([0-9\.\,]+)$', line_clean, flags=re.IGNORECASE)
        if not m:
            continue

        model_part = m.group(1).strip()
        price_part = m.group(2).strip()
        clean_price = re.sub(r'[^0-9]', '', price_part)
        if not clean_price or not model_part:
            continue

        q_norm = normalize_code(model_part)
        found_p = prod_map.get(q_norm)

        if not found_p:
            for p_norm, p in prod_map.items():
                if (q_norm.startswith("rg") and q_norm[2:] == p_norm[2:]) or (q_norm.startswith("dh") and q_norm[2:] == p_norm[2:]):
                    found_p = p
                    break

        if not found_p:
            for p_norm, p in prod_map.items():
                if q_norm in p_norm or (len(q_norm) >= 5 and p_norm in q_norm):
                    found_p = p
                    break

        if found_p:
            found_p["Harga_MD"] = clean_price
            found_p["Harga_ADP"] = clean_price
            found_p["Harga_MDP"] = clean_price
            found_p["Harga_DPP"] = clean_price
            found_p["Harga_IPP"] = clean_price
            found_p["Harga_SDP"] = clean_price
            updated_count += 1
            results_detail.append(f"• {found_p.get('Model', model_part)} : {format_rupiah_num(clean_price)}")
        else:
            brand = infer_brand(model_part)
            new_p = {k: "" for k in PRICELIST_FIELDNAMES}
            new_p["Model"] = model_part
            new_p["Brand"] = brand
            new_p["Description"] = f"{brand} {model_part}"
            new_p["Harga_MD"] = clean_price
            new_p["Harga_ADP"] = clean_price
            new_p["Harga_MDP"] = clean_price
            new_p["Harga_DPP"] = clean_price
            new_p["Harga_IPP"] = clean_price
            new_p["Harga_SDP"] = clean_price
            new_p["Category"] = "Manual Update"
            new_p["Sumber"] = "User Direct Update"
            new_p["Warranty"] = "Official Warranty"
            products.append(new_p)
            prod_map[q_norm] = new_p
            added_count += 1
            results_detail.append(f"• {model_part} (Baru) : {format_rupiah_num(clean_price)}")

    total_modified = updated_count + added_count
    if total_modified == 0:
        return False, "Tidak ditemukan pasangan model dan harga yang valid dalam teks tersebut.", 0

    with open(PRICELIST_PATH, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=PRICELIST_FIELDNAMES)
        writer.writeheader()
        for r in products:
            writer.writerow({k: r.get(k, "") for k in PRICELIST_FIELDNAMES})

    summary = f"✅ Berhasil memperbarui {total_modified} produk ({updated_count} diupdate, {added_count} baru):\n" + "\n".join(results_detail[:25])
    if len(results_detail) > 25:
        summary += f"\n... dan {len(results_detail) - 25} produk lainnya."
    return True, summary, total_modified
