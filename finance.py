from datetime import datetime
from database import get_balance, get_category_breakdown, get_transactions_by_period

def format_rupiah(amount: float) -> str:
    # Format e.g. 150000 -> Rp 150.000
    formatted = f"{int(amount):,}".replace(",", ".")
    return f"Rp {formatted}"

def generate_balance_summary(user_id: int) -> str:
    bal = get_balance(user_id)
    inc = format_rupiah(bal['total_income'])
    exp = format_rupiah(bal['total_expense'])
    sisa = format_rupiah(bal['balance'])
    status_icon = "🟢" if bal['balance'] >= 0 else "🔴"

    return (
        "💰 *SALDO KEUANGAN*\n"
        f"• Pemasukan: `{inc}`\n"
        f"• Pengeluaran: `{exp}`\n"
        "━━━━━━━━━━━━\n"
        f"{status_icon} *Saldo:* `{sisa}`"
    )

def generate_financial_report(user_id: int, period: str = "month") -> str:
    period_label = {
        "today": "Hari Ini",
        "week": "7 Hari",
        "month": "Bulan Ini",
        "all": "Semua Waktu"
    }.get(period, "Bulan Ini")

    summary = get_category_breakdown(user_id, period)
    transactions = get_transactions_by_period(user_id, period)

    tot_inc = format_rupiah(summary['total_income'])
    tot_exp = format_rupiah(summary['total_expense'])
    net = format_rupiah(summary['balance'])

    lines = [
        f"📊 *Laporan ({period_label})*",
        f"• Masuk: `{tot_inc}`",
        f"• Keluar: `{tot_exp}`",
        f"• Bersih: `{net}`"
    ]

    if summary['expense_categories']:
        lines.append("\n📌 *Pengeluaran:*")
        for cat in summary['expense_categories'][:5]:
            lines.append(f"• {cat['category']}: {format_rupiah(cat['total'])}")

    if transactions:
        lines.append("\n🕒 *Transaksi Terakhir:*")
        for t in transactions[:5]:
            icon = "+" if t['type'] == 'income' else "-"
            desc = f" ({t['description']})" if t['description'] else ""
            lines.append(f"• {icon}{format_rupiah(t['amount'])} {t['category']}{desc}")

    return "\n".join(lines)

def export_financial_report_excel(user_id: int, period: str = "all") -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from config import DATA_DIR

    excel_dir = DATA_DIR / "exports"
    excel_dir.mkdir(exist_ok=True)

    summary = get_category_breakdown(user_id, period)
    transactions = get_transactions_by_period(user_id, period)

    wb = Workbook()
    
    # --- Sheet 1: Ringkasan ---
    ws1 = wb.active
    ws1.title = "Ringkasan Keuangan"
    ws1.views.sheetView[0].showGridLines = True

    header_font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
    title_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=11, bold=True)
    regular_font = Font(name="Calibri", size=11)
    
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    income_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    expense_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    ws1.merge_cells('A1:D1')
    ws1['A1'] = "LAPORAN KEUANGAN PRIBADI - SELOBROW"
    ws1['A1'].font = header_font

    ws1['A3'] = "Total Pemasukan"
    ws1['B3'] = summary['total_income']
    ws1['B3'].number_format = '"Rp "#,##0'
    ws1['B3'].font = bold_font

    ws1['A4'] = "Total Pengeluaran"
    ws1['B4'] = summary['total_expense']
    ws1['B4'].number_format = '"Rp "#,##0'
    ws1['B4'].font = bold_font

    ws1['A5'] = "Saldo Bersih"
    ws1['B5'] = summary['balance']
    ws1['B5'].number_format = '"Rp "#,##0'
    ws1['B5'].font = bold_font

    for r in range(3, 6):
        ws1[f'A{r}'].font = regular_font
        ws1[f'A{r}'].border = thin_border
        ws1[f'B{r}'].border = thin_border

    # Rincian Kategori Pengeluaran
    start_row = 7
    ws1.cell(row=start_row, column=1, value="Kategori Pengeluaran").font = title_font
    ws1.cell(row=start_row, column=1).fill = header_fill
    ws1.cell(row=start_row, column=2, value="Total (Rp)").font = title_font
    ws1.cell(row=start_row, column=2).fill = header_fill
    ws1.cell(row=start_row, column=3, value="Frekuensi").font = title_font
    ws1.cell(row=start_row, column=3).fill = header_fill

    cur_row = start_row + 1
    for cat in summary['expense_categories']:
        ws1.cell(row=cur_row, column=1, value=cat['category']).font = regular_font
        c2 = ws1.cell(row=cur_row, column=2, value=cat['total'])
        c2.font = regular_font
        c2.number_format = '"Rp "#,##0'
        c3 = ws1.cell(row=cur_row, column=3, value=cat['count'])
        c3.font = regular_font
        for col in range(1, 4):
            ws1.cell(row=cur_row, column=col).border = thin_border
        cur_row += 1

    # --- Sheet 2: Semua Transaksi ---
    ws2 = wb.create_sheet(title="Daftar Transaksi")
    ws2.views.sheetView[0].showGridLines = True

    headers = ["ID", "Tanggal", "Tipe", "Kategori", "Keterangan", "Nominal (Rp)"]
    for col_num, h_text in enumerate(headers, 1):
        cell = ws2.cell(row=1, column=col_num, value=h_text)
        cell.font = title_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for idx, t in enumerate(transactions, 2):
        ws2.cell(row=idx, column=1, value=t['id']).font = regular_font
        ws2.cell(row=idx, column=2, value=t['date']).font = regular_font
        
        t_type = "Pemasukan" if t['type'] == 'income' else "Pengeluaran"
        c_type = ws2.cell(row=idx, column=3, value=t_type)
        c_type.font = regular_font
        c_type.fill = income_fill if t['type'] == 'income' else expense_fill

        ws2.cell(row=idx, column=4, value=t['category']).font = regular_font
        ws2.cell(row=idx, column=5, value=t['description']).font = regular_font
        
        c_amt = ws2.cell(row=idx, column=6, value=t['amount'])
        c_amt.font = regular_font
        c_amt.number_format = '"Rp "#,##0'
        
        for c in range(1, 7):
            ws2.cell(row=idx, column=c).border = thin_border

    # Auto-adjust column width
    for ws in [ws1, ws2]:
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    export_path = excel_dir / f"Laporan_Keuangan_Selobrow_{user_id}.xlsx"
    wb.save(export_path)
    return str(export_path)

